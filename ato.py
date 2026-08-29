import time
import re
import pandas as pd

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from webdriver_manager.chrome import ChromeDriverManager

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment


# ============================================================
# MESLEK GRUP NO
# ============================================================

KOMITE_NO = input("Ankara için meslek grup no girin: ").strip()

if not KOMITE_NO.isdigit():
    print("HATA: Meslek grup no sadece sayı olmalıdır.")
    exit()

KOMITE_NO = int(KOMITE_NO)


# ============================================================
# AYARLAR
# ============================================================

SEHIR = "Ankara"

URL = (
    f"https://www.atonet.org.tr/"
    f"KomiteListe?komiteKodu={KOMITE_NO}"
)

KOMITE_URL = (
    f"https://www.atonet.org.tr/"
    f"Komite/{KOMITE_NO}"
)

OUTPUT = f"ato_{KOMITE_NO}.xlsx"


# ============================================================
# ANKARA İLÇELERİ
# ============================================================

SEHIR_ILCELERI = {

    "AKYURT",
    "ALTINDAĞ",
    "AYAŞ",
    "BALA",
    "BEYPAZARI",
    "ÇAMLIDERE",
    "ÇANKAYA",
    "ÇUBUK",
    "ELMADAĞ",
    "ETİMESGUT",
    "EVREN",
    "GÖLBAŞI",
    "GÜDÜL",
    "HAYMANA",
    "KAHRAMANKAZAN",
    "KALECİK",
    "KEÇİÖREN",
    "KIZILCAHAMAM",
    "MAMAK",
    "NALLIHAN",
    "POLATLI",
    "PURSAKLAR",
    "SİNCAN",
    "ŞEREFLİKOÇHİSAR",
    "YENİMAHALLE"

}


# ============================================================
# TÜRKÇE KARAKTER NORMALİZASYONU
# ============================================================

def normalize(text):

    text = str(text).strip().upper()

    ceviri = str.maketrans(
        "ÇĞİÖŞÜ",
        "CGIOSU"
    )

    return text.translate(ceviri)


# ============================================================
# ADRESTEN İLÇE BUL
# ============================================================

def ilce_bul(adres):

    if pd.isna(adres):
        return ""

    adres = str(adres).strip()

    if not adres:
        return ""

    adres_normal = normalize(adres)

    sehir_normal = normalize(SEHIR)


    # --------------------------------------------------------
    # ADRESİN SONUNDA /ANKARA VAR MI?
    # --------------------------------------------------------

    if not re.search(
        rf"/\s*{re.escape(sehir_normal)}\s*$",
        adres_normal
    ):
        return ""


    # --------------------------------------------------------
    # SADECE GERÇEK İLÇELERDEN BİRİNİ KABUL ET
    # --------------------------------------------------------

    for ilce in SEHIR_ILCELERI:

        ilce_normal = normalize(ilce)

        pattern = (
            rf"(?:^|[\s,])"
            rf"{re.escape(ilce_normal)}"
            rf"\s*/\s*"
            rf"{re.escape(sehir_normal)}"
            rf"\s*$"
        )

        if re.search(
            pattern,
            adres_normal
        ):

            return ilce


    return ""


# ============================================================
# CHROME
# ============================================================

print()
print("=" * 70)
print(f"ATO {SEHIR.upper()} - MESLEK GRUBU {KOMITE_NO}")
print("=" * 70)

print()
print("Chrome başlatılıyor...")


options = Options()

options.add_argument(
    "--start-maximized"
)


driver = webdriver.Chrome(
    service=Service(
        ChromeDriverManager().install()
    ),
    options=options
)


# ============================================================
# KOMİTE TANIMINI AL
# ============================================================

print()
print("Meslek grup bilgisi alınıyor...")

driver.get(KOMITE_URL)

time.sleep(2)


try:

    komite_baslik = driver.find_element(
        By.CSS_SELECTOR,
        "h5.sayfa-baslik"
    ).text.strip()


    # --------------------------------------------------------
    # "MESLEK KOMİTESİ" İFADESİNİ TEMİZLE
    # --------------------------------------------------------

    MESLEK_GRUP_TANIM = re.sub(
        r"\s+MESLEK KOMİTESİ$",
        "",
        komite_baslik,
        flags=re.IGNORECASE
    ).strip()


    print(
        f"Meslek Grup Tanım: {MESLEK_GRUP_TANIM}"
    )


except Exception as e:

    print()
    print("UYARI: Meslek grup tanımı alınamadı.")
    print(f"Hata: {e}")

    MESLEK_GRUP_TANIM = "yok"


# ============================================================
# FİRMA LİSTESİ SAYFASINI AÇ
# ============================================================

print()
print("ATO firma listesi açılıyor...")

driver.get(URL)

time.sleep(3)


# ============================================================
# TÜM SAYFALARI TARA
# ============================================================

firmalar = []

sayfa = 1


while True:

    print()
    print(
        f"Sayfa {sayfa} okunuyor..."
    )

    time.sleep(1)


    # --------------------------------------------------------
    # TABLO SATIRLARI
    # --------------------------------------------------------

    rows = driver.find_elements(
        By.CSS_SELECTOR,
        "#TBL_KOMITE tbody tr"
    )

    print(
        f"  Satır sayısı: {len(rows)}"
    )


    if not rows:

        print(
            "  TABLO BULUNAMADI!"
        )

        break


    # --------------------------------------------------------
    # FİRMALARI OKU
    # --------------------------------------------------------

    for row in rows:

        cells = row.find_elements(
            By.TAG_NAME,
            "td"
        )


        if len(cells) < 5:
            continue


        sicil = cells[1].text.strip()

        unvan = cells[2].text.strip()

        adres = cells[3].text.strip()

        telefon = cells[4].text.strip()


        # ----------------------------------------------------
        # BOŞ SATIRLARI ATLA
        # ----------------------------------------------------

        if not sicil or not unvan:
            continue


        # ----------------------------------------------------
        # TELEFON
        # ----------------------------------------------------

        if not telefon:
            telefon = "yok"


        # ----------------------------------------------------
        # İLÇE
        # ----------------------------------------------------

        ilce = ilce_bul(
            adres
        )


        # ----------------------------------------------------
        # FİRMAYI EKLE
        # ----------------------------------------------------

        firmalar.append({

            "Ticaret Sicil No":
                sicil,

            "Unvan":
                unvan,

            "Adres":
                adres,

            "İlçe":
                ilce,

            "Web":
                "",

            "Mail":
                "",

            "Firma Tel":
                telefon,

            "Meslek Grup No":
                KOMITE_NO,

            "Meslek Grup Tanım":
                MESLEK_GRUP_TANIM

        })


    # ========================================================
    # SONRAKİ BUTONU BUL
    # ========================================================

    try:

        next_button = driver.find_element(
            By.CSS_SELECTOR,
            "#TBL_KOMITE_next"
        )

    except Exception:

        print()
        print(
            "  Sonraki sayfa butonu bulunamadı."
        )

        break


    # --------------------------------------------------------
    # SON SAYFA MI?
    # --------------------------------------------------------

    class_name = (
        next_button
        .get_attribute("class")
        or ""
    )


    if "disabled" in class_name:

        print()
        print(
            "  Son sayfaya ulaşıldı."
        )

        break


    # --------------------------------------------------------
    # ŞU ANKİ İLK FİRMAYI HATIRLA
    # --------------------------------------------------------

    eski_ilk_sicil = ""

    try:

        eski_ilk_sicil = (
            rows[0]
            .find_elements(
                By.TAG_NAME,
                "td"
            )[1]
            .text
            .strip()
        )

    except Exception:

        pass


    # ========================================================
    # SONRAKİ SAYFAYA TIKLA
    # ========================================================

    driver.execute_script(
        "arguments[0].click();",
        next_button
    )


    # ========================================================
    # TABLONUN DEĞİŞMESİNİ BEKLE
    # ========================================================

    baslangic = time.time()


    while time.time() - baslangic < 10:

        time.sleep(0.2)


        yeni_rows = driver.find_elements(
            By.CSS_SELECTOR,
            "#TBL_KOMITE tbody tr"
        )


        if not yeni_rows:
            continue


        try:

            yeni_ilk_sicil = (
                yeni_rows[0]
                .find_elements(
                    By.TAG_NAME,
                    "td"
                )[1]
                .text
                .strip()
            )

        except Exception:

            continue


        if (
            yeni_ilk_sicil
            and yeni_ilk_sicil
            != eski_ilk_sicil
        ):

            break


    sayfa += 1


# ============================================================
# BROWSER KAPAT
# ============================================================

driver.quit()


# ============================================================
# DATAFRAME
# ============================================================

df = pd.DataFrame(
    firmalar
)


# ============================================================
# AYNI FİRMALARI TEMİZLE
# ============================================================

if not df.empty:

    df = df.drop_duplicates(
        subset=[
            "Ticaret Sicil No"
        ],
        keep="first"
    )


# ============================================================
# EXCEL
# ============================================================

df.to_excel(
    OUTPUT,
    index=False
)


# ============================================================
# EXCEL BİÇİMLENDİRME
# ============================================================

print()
print("Excel biçimlendiriliyor...")


wb = load_workbook(
    OUTPUT
)

ws = wb.active


# ------------------------------------------------------------
# TÜM HÜCRELER
# Arial - 11 punto
# ------------------------------------------------------------

for row in ws.iter_rows():

    for cell in row:

        cell.font = Font(
            name="Arial",
            size=11
        )

        cell.alignment = Alignment(
            vertical="center"
        )


# ------------------------------------------------------------
# BAŞLIK SATIRI
# Beyaz + Bold
# Dolgu: #46bdc6
# ------------------------------------------------------------

header_fill = PatternFill(
    fill_type="solid",
    fgColor="46BDC6"
)


for cell in ws[1]:

    cell.font = Font(
        name="Arial",
        size=11,
        bold=True,
        color="FFFFFF"
    )

    cell.fill = header_fill

    cell.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )


# ------------------------------------------------------------
# SÜTUN GENİŞLİKLERİ
# ------------------------------------------------------------

for column in ws.columns:

    max_length = 0

    column_letter = column[0].column_letter


    for cell in column:

        if cell.value is not None:

            cell_length = len(
                str(cell.value)
            )

            if cell_length > max_length:

                max_length = cell_length


    # Çok uzun sütunların aşırı genişlemesini önle
    adjusted_width = min(
        max_length + 2,
        60
    )


    ws.column_dimensions[
        column_letter
    ].width = adjusted_width


# ------------------------------------------------------------
# BAŞLIK SATIRI YÜKSEKLİĞİ
# ------------------------------------------------------------

ws.row_dimensions[1].height = 22


# ============================================================
# EXCEL KAYDET
# ============================================================

wb.save(
    OUTPUT
)


# ============================================================
# İSTATİSTİK
# ============================================================

toplam = len(df)


if toplam > 0:

    ilce_bulunan = (
        df["İlçe"]
        .fillna("")
        .astype(str)
        .str.strip()
        .ne("")
        .sum()
    )


    ilce_bulunamayan = (
        toplam
        - ilce_bulunan
    )


    telefon_var = (
        df["Firma Tel"]
        .fillna("")
        .astype(str)
        .str.strip()
        .ne("yok")
        .sum()
    )


    telefon_yok = (
        df["Firma Tel"]
        .fillna("")
        .astype(str)
        .str.strip()
        .eq("yok")
        .sum()
    )

else:

    ilce_bulunan = 0
    ilce_bulunamayan = 0
    telefon_var = 0
    telefon_yok = 0


# ============================================================
# SONUÇ
# ============================================================

print()
print("=" * 70)
print("TAMAMLANDI")
print("=" * 70)

print(
    f"Meslek Grup No    : {KOMITE_NO}"
)

print(
    f"Meslek Grup Tanım : {MESLEK_GRUP_TANIM}"
)

print(
    f"Toplam firma      : {toplam}"
)

print(
    f"İlçe bulunan      : {ilce_bulunan}"
)

print(
    f"İlçe bulunamayan  : {ilce_bulunamayan}"
)

print(
    f"Telefon bulunan   : {telefon_var}"
)

print(
    f"Telefon bulunamayan: {telefon_yok}"
)

print(
    f"Excel             : {OUTPUT}"
)

print("=" * 70)
